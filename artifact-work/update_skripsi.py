from __future__ import annotations

import math
from copy import deepcopy
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH


SOURCE = Path(r"D:\skripsi\Skripsi_GIS_revisi_halaman_pengesahan.docx")
OUTPUT = Path(r"D:\projek UMKM\gis-umkm-mapping-main\Skripsi_GIS_diperbarui_sesuai_proyek_pembaruan.docx")
DIAGRAM_DIR = Path(r"D:\projek UMKM\gis-umkm-mapping-main\artifact-work\updated-diagrams")
FONT_REGULAR = Path(r"C:\Windows\Fonts\arial.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\arialbd.ttf")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT_REGULAR), size)


def centered_text(
    draw: ImageDraw.ImageDraw,
    bounds: tuple[int, int, int, int],
    text: str,
    size: int,
    bold: bool = False,
    spacing: int = 5,
) -> None:
    text_font = font(size, bold)
    box = draw.multiline_textbbox((0, 0), text, font=text_font, spacing=spacing, align="center")
    width = box[2] - box[0]
    height = box[3] - box[1]
    x1, y1, x2, y2 = bounds
    draw.multiline_text(
        ((x1 + x2 - width) / 2, (y1 + y2 - height) / 2 - box[1]),
        text,
        font=text_font,
        fill="black",
        spacing=spacing,
        align="center",
    )


def box(
    draw: ImageDraw.ImageDraw,
    bounds: tuple[int, int, int, int],
    text: str,
    size: int = 28,
    radius: int = 18,
    width: int = 3,
) -> None:
    draw.rounded_rectangle(bounds, radius=radius, fill="white", outline="black", width=width)
    centered_text(draw, bounds, text, size)


def arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    width: int = 4,
    head: int = 14,
) -> None:
    draw.line((start, end), fill="black", width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    left = (
        end[0] - head * math.cos(angle - math.pi / 6),
        end[1] - head * math.sin(angle - math.pi / 6),
    )
    right = (
        end[0] - head * math.cos(angle + math.pi / 6),
        end[1] - head * math.sin(angle + math.pi / 6),
    )
    draw.polygon((end, left, right), fill="black")


def title(draw: ImageDraw.ImageDraw, width: int, text: str, size: int = 38) -> None:
    centered_text(draw, (30, 15, width - 30, 80), text, size, True)


def save_image(name: str, size: tuple[int, int], paint) -> Path:
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    paint(draw, *size)
    path = DIAGRAM_DIR / name
    image.save(path, format="PNG", optimize=True)
    return path


def research_flow(draw: ImageDraw.ImageDraw, width: int, height: int) -> None:
    title(draw, width, "Alur Penelitian dan Pengembangan", 36)
    labels = [
        "Identifikasi masalah dan kebutuhan data spasial",
        "Inventarisasi dan pembersihan data UMKM",
        "Pemberian status kualitas dan koordinat lokasi",
        "Perancangan SIG, pengajuan publik, dan basis data",
        "Implementasi K-Means++ dan jarak Haversine",
        "Implementasi admin, verifikasi, pelacakan, dan storage",
        "Pengujian build, fungsi, keamanan, dan hasil klaster",
        "Analisis hasil, kesimpulan, dan saran",
    ]
    left, right = 180, width - 180
    top, box_height, gap = 105, 94, 35
    for index, label in enumerate(labels):
        y1 = top + index * (box_height + gap)
        y2 = y1 + box_height
        box(draw, (left, y1, right, y2), label, 24)
        if index < len(labels) - 1:
            arrow(draw, (width // 2, y2), (width // 2, y2 + gap - 5), 4, 13)


def architecture(draw: ImageDraw.ImageDraw, width: int, height: int) -> None:
    title(draw, width, "Arsitektur Sistem Zonasi UMKM", 38)
    public = (55, 135, 335, 280)
    admin = (55, 555, 335, 700)
    react = (470, 105, 820, 260)
    kmeans = (470, 340, 820, 490)
    service = (470, 580, 820, 730)
    leaflet = (1010, 105, 1345, 260)
    supabase = (990, 330, 1365, 500)
    storage = (1010, 580, 1345, 730)
    box(draw, public, "Pengguna Publik\nBrowser / Mobile", 27)
    box(draw, admin, "Admin / Verifikator\nViewer", 27)
    box(draw, react, "React + Vite\nPeta dan pencarian\nPengajuan dan pelacakan", 22)
    box(draw, kmeans, "Utility K-Means\nHaversine + WCSS", 27)
    box(draw, service, "Service Layer\nSupabase JS + RPC", 27)
    box(draw, leaflet, "Leaflet\nPeta + Marker Cluster", 27)
    box(draw, supabase, "Supabase\nAuth + PostgreSQL\nRLS + Trigger", 27)
    box(draw, storage, "Supabase Storage\nFoto pengajuan privat", 27)
    arrow(draw, (public[2], 205), (react[0], 205))
    arrow(draw, (admin[2], 625), (service[0], 655))
    arrow(draw, (react[2], 180), (leaflet[0], 180))
    arrow(draw, (645, react[3]), (645, kmeans[1]))
    arrow(draw, (645, kmeans[3]), (645, service[1]))
    arrow(draw, (service[2], 635), (supabase[0], 435))
    arrow(draw, (service[2], 685), (storage[0], 655))
    arrow(draw, (supabase[0], 385), (react[2], 230))


def data_flow(draw: ImageDraw.ImageDraw, width: int, height: int) -> None:
    draw.text((30, 18), "ALUR PUBLIK", font=font(24, True), fill="black")
    draw.text((30, 235), "ALUR ADMIN", font=font(24, True), fill="black")
    public_boxes = [
        ((50, 70, 300, 175), "Pengguna publik"),
        ((370, 55, 700, 190), "React\nPeta + formulir"),
        ((770, 45, 1110, 195), "Service layer\nvalidasi + upload"),
        ((1190, 45, 1560, 195), "Supabase\nPostgreSQL + Storage"),
        ((1640, 45, 2010, 195), "Peta, daftar, status\ndan kode pelacakan"),
        ((2090, 70, 2420, 175), "Share, rute,\nUMKM terdekat"),
    ]
    admin_boxes = [
        ((50, 285, 300, 390), "Admin berautentikasi"),
        ((370, 270, 700, 405), "Dashboard + CRUD\nVerifikasi + K-Means"),
        ((770, 260, 1110, 410), "Pengajuan + pemulihan\nImpor + audit"),
        ((1190, 260, 1560, 410), "RPC + RLS + trigger\ntransaksi atomik"),
        ((1640, 270, 2010, 405), "UMKM + riwayat\nK-Means + audit"),
        ((2090, 285, 2420, 390), "Data publik\nterbarui"),
    ]
    for bounds, label in public_boxes + admin_boxes:
        box(draw, bounds, label, 22, 16, 3)
    for row in (public_boxes, admin_boxes):
        for current, following in zip(row, row[1:]):
            arrow(draw, (current[0][2], (current[0][1] + current[0][3]) // 2), (following[0][0], (following[0][1] + following[0][3]) // 2), 3, 12)
    arrow(draw, (1375, 260), (1375, 195), 3, 12)


def erd(draw: ImageDraw.ImageDraw, width: int, height: int) -> None:
    title(draw, width, "Entity Relationship Diagram Ringkas", 36)
    entities = {
        "profiles": ((560, 100, 940, 275), "ADMIN_PROFILES\nuser_id PK/FK\nrole, is_active"),
        "umkm": ((45, 310, 405, 560), "UMKM\nid PK\natribut usaha\nanalysis/display lat-lng\nlocation_accuracy\nverified_by FK"),
        "runs": ((1095, 310, 1455, 560), "KMEANS_RUNS\nid PK\nk_value, wcss\ncentroids, cluster_stats\ndataset_hash\ninput_snapshot"),
        "submissions": ((455, 350, 1045, 595), "UMKM_SUBMISSIONS\nid PK, tracking_code\ndata usaha + koordinat opsional\nphoto_path + photo_kind\nstatus + review_note\nreviewed_by / approved_umkm_id FK"),
        "audit": ((45, 680, 405, 895), "AUDIT_LOGS\nid PK\nrecord_id, action\nold_data, new_data\nactor_id"),
        "imports": ((455, 680, 745, 895), "IMPORT_BATCHES\nid PK\nsource_name\njumlah baris\nstatus, imported_by"),
        "recovery": ((795, 680, 1155, 895), "RECOVERY_REQUESTS\nid PK\ndata pemohon\nstatus\nhandled_by / handled_at"),
        "storage": ((1205, 680, 1455, 895), "STORAGE OBJECTS\nbucket privat\nfoto pengajuan\nsigned URL admin"),
    }
    for bounds, label in entities.values():
        box(draw, bounds, label, 22, 14, 3)
    arrow(draw, (640, 275), (355, 310), 3, 11)
    arrow(draw, (860, 275), (1145, 310), 3, 11)
    arrow(draw, (750, 275), (750, 350), 3, 11)
    arrow(draw, (455, 470), (405, 470), 3, 11)
    arrow(draw, (1045, 500), (1095, 500), 3, 11)
    arrow(draw, (250, 560), (250, 680), 3, 11)
    arrow(draw, (600, 595), (600, 680), 3, 11)
    arrow(draw, (900, 595), (965, 680), 3, 11)
    arrow(draw, (1010, 500), (1330, 680), 3, 11)


def actor(draw: ImageDraw.ImageDraw, center_x: int, top: int, label: str) -> None:
    draw.ellipse((center_x - 28, top, center_x + 28, top + 56), outline="black", width=3)
    draw.line((center_x, top + 56, center_x, top + 165), fill="black", width=3)
    draw.line((center_x - 55, top + 95, center_x + 55, top + 95), fill="black", width=3)
    draw.line((center_x, top + 165, center_x - 48, top + 235), fill="black", width=3)
    draw.line((center_x, top + 165, center_x + 48, top + 235), fill="black", width=3)
    centered_text(draw, (center_x - 140, top + 245, center_x + 140, top + 305), label, 25)


def ellipse(draw: ImageDraw.ImageDraw, bounds: tuple[int, int, int, int], text: str, size: int = 23) -> None:
    draw.ellipse(bounds, fill="white", outline="black", width=3)
    centered_text(draw, bounds, text, size)


def use_case(draw: ImageDraw.ImageDraw, width: int, height: int) -> None:
    title(draw, width, "Use Case Sistem", 38)
    actor(draw, 115, 220, "Pengguna Publik")
    actor(draw, 1285, 220, "Administrator")
    public_cases = [
        ((300, 105, 650, 205), "Melihat peta, daftar,\ndan zonasi"),
        ((300, 255, 650, 355), "Mencari, memfilter,\ndan memilih UMKM"),
        ((300, 405, 650, 505), "UMKM terdekat, rute,\ndan berbagi tautan"),
        ((300, 555, 650, 655), "Mengajukan UMKM\ndan mengunggah foto"),
        ((300, 705, 650, 805), "Melacak status dan\nmeminta pemulihan kode"),
    ]
    admin_cases = [
        ((750, 105, 1100, 205), "Dashboard dan\nkualitas data"),
        ((750, 255, 1100, 355), "Mengelola data dan\nmemverifikasi lokasi"),
        ((750, 405, 1100, 505), "Meninjau pengajuan dan\npemulihan kode"),
        ((750, 555, 1100, 655), "Menjalankan dan\nmenyimpan K-Means"),
        ((750, 705, 1100, 805), "Impor, ekspor,\ndan melihat audit"),
    ]
    for bounds, label in public_cases + admin_cases:
        ellipse(draw, bounds, label)
    for bounds, _ in public_cases:
        draw.line((170, 315, bounds[0], (bounds[1] + bounds[3]) // 2), fill="black", width=2)
    for bounds, _ in admin_cases:
        draw.line((1230, 315, bounds[2], (bounds[1] + bounds[3]) // 2), fill="black", width=2)


def location_flow(draw: ImageDraw.ImageDraw, width: int, height: int) -> None:
    nodes = [
        ((40, 80, 320, 210), "Data lama atau\npengajuan publik"),
        ((400, 80, 710, 210), "Normalisasi dan\nkoordinat awal"),
        ((790, 80, 1110, 210), "Status kualitas\nlokasi"),
        ((1210, 80, 1570, 210), "Admin / Verifikator\nmemeriksa lokasi"),
        ((1210, 310, 1570, 440), "tepat\nverified_at + verified_by"),
        ((790, 310, 1110, 440), "perkiraan_kecamatan\natau tepat"),
        ((400, 310, 710, 440), "Koordinat analisis\nanalysis_lat/lng"),
        ((40, 310, 320, 440), "K-Means\nHaversine + WCSS"),
        ((790, 560, 1110, 690), "belum_terverifikasi\ndikecualikan"),
        ((400, 560, 710, 690), "Koordinat tampilan\ndisplay_lat/lng"),
        ((40, 560, 320, 690), "Peta, rute, dan\nUMKM terdekat"),
        ((1210, 560, 1570, 690), "Statistik kualitas\ndan audit"),
    ]
    for bounds, label in nodes:
        box(draw, bounds, label, 26, 16, 3)
    arrow(draw, (320, 145), (400, 145), 3, 13)
    arrow(draw, (710, 145), (790, 145), 3, 13)
    arrow(draw, (1110, 145), (1210, 145), 3, 13)
    arrow(draw, (1390, 210), (1390, 310), 3, 13)
    arrow(draw, (1210, 375), (1110, 375), 3, 13)
    arrow(draw, (790, 375), (710, 375), 3, 13)
    arrow(draw, (400, 375), (320, 375), 3, 13)
    arrow(draw, (950, 210), (950, 310), 3, 13)
    arrow(draw, (950, 210), (950, 560), 3, 13)
    arrow(draw, (555, 210), (555, 560), 3, 13)
    arrow(draw, (400, 625), (320, 625), 3, 13)
    arrow(draw, (1110, 625), (1210, 625), 3, 13)


def kmeans_flow(draw: ImageDraw.ImageDraw, width: int, height: int) -> None:
    title(draw, width, "Flowchart Proses K-Means Spasial", 36)
    labels = [
        "Ambil UMKM aktif, published, dan mappable",
        "Urutkan dan bentuk snapshot id-lat-lng",
        "Tentukan K dan seed deterministik",
        "Inisialisasi centroid dengan K-Means++",
        "Alokasikan titik memakai jarak Haversine",
        "Hitung centroid baru; ulangi sampai konvergen atau 100 iterasi",
        "Hitung WCSS, radius, dan statistik klaster",
        "Simpan parameter, hash, snapshot, dan hasil run",
    ]
    left, right = 150, width - 150
    top, box_height, gap = 100, 86, 34
    for index, label in enumerate(labels):
        y1 = top + index * (box_height + gap)
        y2 = y1 + box_height
        box(draw, (left, y1, right, y2), label, 23, 16, 3)
        if index < len(labels) - 1:
            arrow(draw, (width // 2, y2), (width // 2, y2 + gap - 5), 3, 12)


def test_flow(draw: ImageDraw.ImageDraw, width: int, height: int) -> None:
    top_nodes = [
        ((40, 10, 360, 115), "Siapkan lingkungan\ndan data uji"),
        ((440, 10, 760, 115), "Jalankan skenario\nblack-box"),
        ((840, 10, 1160, 115), "Catat input, output,\nstatus, dan bukti"),
        ((1240, 10, 1560, 115), "Uji K-Means berulang\npada beberapa K"),
    ]
    bottom_nodes = [
        ((1240, 185, 1560, 290), "Analisis WCSS, iterasi,\ndan komposisi klaster"),
        ((840, 185, 1160, 290), "Evaluasi kegagalan\ndan perbaiki"),
        ((440, 185, 760, 290), "Jalankan ulang\npengujian terkait"),
        ((40, 185, 360, 290), "Dokumentasikan\nhasil final"),
    ]
    for bounds, label in top_nodes + bottom_nodes:
        box(draw, bounds, label, 24, 16, 3)
    for current, following in zip(top_nodes, top_nodes[1:]):
        arrow(draw, (current[0][2], 62), (following[0][0], 62), 3, 13)
    arrow(draw, (1400, 115), (1400, 185), 3, 13)
    for current, following in zip(bottom_nodes, bottom_nodes[1:]):
        arrow(draw, (current[0][0], 237), (following[0][2], 237), 3, 13)


def diagram_files() -> dict[str, Path]:
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    return {
        "image2.png": save_image("image2.png", (1200, 1150), research_flow),
        "image3.png": save_image("image3.png", (1400, 850), architecture),
        "image4.png": save_image("image4.png", (2474, 434), data_flow),
        "image5.png": save_image("image5.png", (1500, 950), erd),
        "image6.png": save_image("image6.png", (1400, 850), use_case),
        "image7.png": save_image("image7.png", (1800, 760), location_flow),
        "image8.png": save_image("image8.png", (1250, 1200), kmeans_flow),
        "image9.png": save_image("image9.png", (1600, 320), test_flow),
    }


def replace_paragraphs(document: Document) -> None:
    replacements = {
        43: "Data Usaha Mikro, Kecil, dan Menengah (UMKM) yang tersebar pada wilayah yang luas sulit dimanfaatkan untuk analisis kewilayahan apabila hanya disajikan dalam bentuk tabel. Penelitian ini mengembangkan sistem informasi geografis berbasis web untuk memetakan UMKM di Provinsi Sulawesi Utara dan membentuk zonasi spasial menggunakan K-Means. Aplikasi menggunakan React 19, Vite 8, Leaflet, Supabase PostgreSQL, autentikasi, Row Level Security, dan penyimpanan objek privat. Layanan publik mencakup pencarian, filter, koleksi kategori, UMKM terdekat, rute, berbagi tautan, pengajuan UMKM dengan foto dan titik lokasi opsional, pelacakan status, serta permintaan pemulihan kode. Panel admin menyediakan pengelolaan data, filter kualitas, verifikasi koordinat, peninjauan pengajuan, impor dan ekspor CSV, audit, serta penyimpanan hasil analisis. Klasterisasi menerapkan K-Means++ deterministik, jarak Haversine, WCSS, radius klaster, snapshot input, dan hash SHA-256. Snapshot JSON proyek tanggal 10 September 2026 memuat 1.802 record; 1.743 lokasi perkiraan dapat dianalisis dan 59 lokasi belum terverifikasi dikecualikan. Eksekusi K=2 sampai K=10 menghasilkan kurva WCSS yang tidak monoton pada K=9, sehingga nilai K optimal belum dapat ditetapkan hanya dari elbow pada satu inisialisasi per K. Pemeriksaan teknis menunjukkan lint dan build produksi berhasil, sedangkan pengujian integrasi Supabase dan black-box lintas perangkat tetap perlu didokumentasikan pada lingkungan final.",
        44: "Kata kunci: UMKM, Sistem Informasi Geografis, K-Means, Haversine, pengajuan publik, kualitas data spasial, Supabase.",
        48: "Micro, Small, and Medium Enterprise (MSME) data distributed over a wide area are difficult to use for regional analysis when presented only as tables. This research develops a web-based geographic information system for mapping MSMEs in North Sulawesi Province and producing spatial zones with K-Means. The application uses React 19, Vite 8, Leaflet, Supabase PostgreSQL, authentication, Row Level Security, and private object storage. Public services include search, filters, curated category collections, nearby MSMEs, directions, link sharing, MSME submissions with an optional photo and location, status tracking, and tracking-code recovery requests. The administration panel supports data management, quality filters, coordinate verification, submission review, CSV import and export, audit history, and stored analysis runs. Clustering uses deterministic K-Means++ initialization, Haversine distance, WCSS, cluster radii, input snapshots, and SHA-256 dataset hashes. The local project snapshot dated September 10, 2026 contains 1,802 records: 1,743 district-estimated locations are analyzable and 59 unverified locations are excluded. Runs for K=2 through K=10 produced a non-monotonic WCSS sequence at K=9, so an optimal K cannot be established from a single-seed elbow curve alone. Linting and the production build completed successfully, while final Supabase integration and cross-device black-box tests still require documented execution in the deployment environment.",
        49: "Keywords: MSME, Geographic Information System, K-Means, Haversine, public submission, spatial data quality, Supabase.",
        52: "Puji syukur dipanjatkan ke hadirat Tuhan Yang Maha Esa atas rahmat dan karunia-Nya sehingga skripsi berjudul “Implementasi Algoritma K-Means Clustering pada Web GIS untuk Pemetaan dan Analisis Zonasi Persebaran UMKM di Sulawesi Utara” dapat disusun. Penelitian ini berfokus pada pemetaan UMKM, layanan pengajuan publik, pengelolaan kualitas koordinat, dan analisis zonasi berbasis klaster spasial.",
        53: "Penulis menyadari bahwa tanpa bantuan dan bimbingan dari berbagai pihak, skripsi ini sulit diselesaikan. Oleh karena itu, penulis menyampaikan ucapan terima kasih kepada:\n1. Direktur Politeknik Negeri Manado;\n2. Ketua Jurusan Teknik Elektro;\n3. Koordinator Program Studi D-IV Teknik Informatika;\n4. Ketua Pelaksana Ujian Skripsi beserta seluruh panitia;\n5. Dosen Pembimbing I dan Dosen Pembimbing II atas arahan, bimbingan, dan masukan selama penyusunan skripsi ini;\n6. Seluruh dosen dan staf pengajar Program Studi D-IV Teknik Informatika atas ilmu, arahan, dan dukungan akademik;\n7. YBLI Yayasan Bina Lentera Insan selaku mitra penelitian yang telah mendukung pelaksanaan penelitian ini;\n8. Orang tua, keluarga, sahabat, dan rekan-rekan yang telah membantu penulis menyelesaikan skripsi ini.",
        54: "Penulis menyadari bahwa penelitian ini masih dapat disempurnakan melalui verifikasi koordinat lapangan, pengujian integrasi pada lingkungan produksi, dan evaluasi pengguna. Kritik dan saran yang membangun sangat diharapkan.",
        68: "2.3 Web Mapping, Layanan Publik, dan Leaflet\t12",
        99: "4.2 Implementasi Layanan Publik dan Peta\t44",
        116: "Lampiran 2. Parameter Implementasi K-Means\t68",
        117: "Lampiran 3. Contoh Format Bukti Pengujian\t69",
        118: "",
        152: "Tabel 4.7 Hasil Eksekusi K-Means pada Snapshot Data Lokal\t54",
        164: "Berdasarkan kondisi tersebut, penelitian ini mengembangkan sistem informasi geografis berbasis web untuk pemetaan dan klasterisasi spasial UMKM di Provinsi Sulawesi Utara. Kontribusi terapannya mengintegrasikan visualisasi peta, layanan pengajuan dan pelacakan publik, administrasi data, kontrol kualitas lokasi, audit perubahan, impor dan ekspor, serta analisis K-Means dalam satu aplikasi.",
        172: "Dari sisi rekayasa perangkat lunak, aplikasi dibangun sebagai sistem web dengan pengalaman publik dan ruang administrasi. Pengguna publik dapat mencari, memfilter, melihat koleksi kategori, menemukan lima UMKM terdekat, membuka rute, membagikan tautan, mengajukan data usaha beserta foto dan titik opsional, melacak status, serta meminta pemulihan kode. Pengguna administrasi mengelola data, meninjau pengajuan, memproses pemulihan kode, memverifikasi lokasi, menjalankan analisis, mengimpor atau mengekspor dataset, dan memeriksa audit. Supabase menyediakan autentikasi, PostgreSQL, RPC, RLS, dan penyimpanan foto privat, sedangkan React/Vite dan Leaflet membentuk lapisan antarmuka.",
        174: "Berdasarkan uraian tersebut, kebutuhan penelitian mencakup representasi spasial, partisipasi publik, integritas data, dan analisis. Peta dan fitur penemuan menyajikan data; alur pengajuan memperbarui data melalui tinjauan; status lokasi dan kontrol akses menjaga kualitas; sedangkan K-Means merangkum pola kedekatan pada titik yang memenuhi syarat.",
        176: "1. Bagaimana merancang dan membangun SIG berbasis web yang menampilkan persebaran UMKM serta mendukung pencarian, filter, UMKM terdekat, rute, dan berbagi tautan?",
        177: "2. Bagaimana menyediakan pengajuan UMKM, foto, kode pelacakan, pemulihan kode, dan peninjauan admin dengan pembatasan akses yang sesuai?",
        178: "3. Bagaimana mengelola perbedaan antara koordinat perkiraan, belum terverifikasi, dan koordinat tepat yang telah diperiksa?",
        179: "4. Bagaimana menerapkan K-Means++ deterministik dengan jarak Haversine serta menyimpan WCSS, centroid, radius, hash, dan snapshot input?",
        180: "5. Bagaimana memverifikasi kualitas implementasi melalui lint, build produksi, eksekusi algoritma, dan skenario black-box pada lingkungan final?",
        183: "1. Mengembangkan aplikasi SIG berbasis web untuk menampilkan dan menemukan data UMKM secara interaktif.",
        184: "2. Mengimplementasikan pengajuan publik, unggahan foto, pelacakan status, pemulihan kode, dan peninjauan admin yang terkontrol.",
        185: "3. Mengimplementasikan klasifikasi akurasi lokasi, verifikasi koordinat, dan K-Means++ dengan jarak Haversine.",
        186: "4. Menyimpan WCSS, centroid, radius, parameter, snapshot, dan hash dataset untuk mendukung evaluasi yang dapat ditelusuri.",
        187: "5. Melaksanakan pemeriksaan statis, build produksi, pengujian algoritma, dan skenario black-box sistem final.",
        190: "1. Bagi pengelola data UMKM, sistem membantu melihat kualitas data, memprioritaskan pekerjaan, meninjau pengajuan, dan memperbaiki koordinat secara bertahap.",
        191: "2. Bagi pelaku UMKM dan masyarakat, sistem menyediakan penemuan usaha, rute, pengajuan data, pelacakan status, serta kanal pemulihan kode.",
        193: "4. Bagi pengembang, implementasi menyediakan pola arsitektur yang memisahkan data publik, pengajuan, administrasi, verifikasi, penyimpanan privat, dan analisis.",
        203: "Batasan penelitian mempertahankan fokus clustering pada koordinat. Atribut omzet, tenaga kerja, modal, digitalisasi, dan kapasitas produksi tidak tersedia secara konsisten sehingga tidak digunakan sebagai dimensi. Foto dan kontak pada pengajuan berfungsi untuk peninjauan administratif, bukan sebagai fitur K-Means, dan aksesnya dibatasi oleh kebijakan basis data serta penyimpanan.",
        206: "BAB I membahas masalah, tujuan, manfaat, batasan, dan sistematika. BAB II memuat teori SIG, layanan peta web, kualitas data lokasi, K-Means++, Haversine, WCSS, teknologi, keamanan, dan reproduksibilitas. BAB III menjelaskan kebutuhan, arsitektur, basis data, pengajuan publik, verifikasi, analisis, dan pengujian. BAB IV membahas implementasi proyek, hasil pemeriksaan teknis, statistik dataset, serta eksekusi K-Means. BAB V menyajikan kesimpulan dan saran.",
        220: "2.3 Web Mapping, Layanan Publik, dan Leaflet",
        224: "Komponen publik mendukung pencarian nama, merek, pemilik, alamat, dan kategori; filter kategori serta koleksi komunitas; pemilihan UMKM melalui parameter URL; berbagi tautan; rute; dan pencarian lima UMKM terdekat dari lokasi browser. Marker, rute, dan jarak publik memakai koordinat tampilan agar konsisten dengan titik yang dilihat pengguna, sedangkan K-Means tetap memakai koordinat analisis.",
        225: "Layanan publik juga menyediakan formulir pengajuan UMKM. Pemohon mengisi identitas usaha, pemilik, WhatsApp, kategori, alamat, catatan, foto opsional, serta titik opsional yang dapat dipilih melalui peta atau pencarian Nominatim OpenStreetMap. Setelah pengajuan tersimpan, sistem memberikan UUID pelacakan untuk memeriksa status tanpa membuka seluruh data pengajuan. Jika kode hilang, pemohon dapat mengirim permintaan pemulihan yang diproses admin.",
        255: "Repositori menggunakan React 19 dan Vite 8 untuk antarmuka, Leaflet 1.9 dan react-leaflet 5 untuk peta, serta Supabase JS 2.112.3 untuk PostgreSQL, autentikasi, RPC, dan Storage. Migrasi mendefinisikan data UMKM, profil admin, hasil K-Means, batch impor, audit, pengajuan publik, foto privat, kode pelacakan, dan permintaan pemulihan.",
        259: "Supabase menjadi backend-as-a-service untuk autentikasi, PostgreSQL, dan object storage. Service layer memuat data publik dan admin, mengeksekusi RPC verifikasi, impor, review pengajuan, serta status pelacakan, dan menghasilkan signed URL foto bagi admin. Pemisahan ini mengurangi keterikatan komponen UI pada detail kueri sekaligus memusatkan validasi batas layanan.",
        260: "PostgreSQL menegakkan pasangan koordinat, rentang nilai, metadata lokasi tepat, konsistensi snapshot K-Means, state pengajuan, dan relasi hasil persetujuan. RLS membatasi data admin dan pengajuan, sementara bucket foto bersifat privat; publik hanya dapat mengunggah ke folder submissions dan admin berwenang memperoleh URL sementara untuk membaca foto.",
        274: "Celah terapan penelitian adalah integrasi peta publik, fitur penemuan, partisipasi pelaku usaha, kualitas koordinat bertingkat, verifikasi, kontrol akses, audit, dan K-Means spasial dalam satu workflow. Kode pelacakan dan RPC status memberi transparansi kepada pemohon tanpa memberikan akses baca langsung ke tabel pengajuan.",
        275: "",
        276: "Kontribusi penelitian bersifat rekayasa sistem dan tata kelola data. Sistem tidak menawarkan algoritma clustering baru, tetapi menyatukan visualisasi, pengajuan dan peninjauan, penyimpanan foto privat, verifikasi lokasi, RBAC, impor atomik, audit perubahan, serta bukti input analisis pada konteks UMKM Sulawesi Utara.",
        282: "Objek teknis penelitian adalah repositori gis-umkm-mapping beserta dataset, antarmuka, service, dan migrasi Supabase. Unit clustering adalah record aktif, dipublikasikan, dan mappable. Unit pengujian mencakup peta dan fitur penemuan, pengajuan dan pelacakan, dashboard admin, CRUD, review, verifikasi, impor dan ekspor, audit, serta K-Means.",
        290: "Tahapan penelitian menghasilkan artefak terukur: kebutuhan dan aturan data; arsitektur; aplikasi publik dan admin; migrasi keamanan; peningkatan kualitas lokasi; pemeriksaan lint/build; eksekusi algoritma; pengujian black-box; serta dokumentasi hasil. Temuan pengujian digunakan untuk memperbaiki sistem sebelum snapshot final ditetapkan.",
        291: "Sumber bukti terdiri atas kode, migrasi SQL, dataset, hasil lint dan build, keluaran K-Means, pengujian aplikasi, screenshot, serta riwayat analisis. Kode membuktikan mekanisme, sedangkan hasil eksekusi membuktikan perilaku. Perbedaan ini mencegah klaim lulus hanya karena fungsi ditemukan pada repositori.",
        294: "Pengguna publik memerlukan peta, daftar, pencarian, filter, koleksi, UMKM terdekat, rute, berbagi tautan, pengajuan, pelacakan, dan pemulihan kode. Admin mengelola data, kualitas, pengajuan, impor, ekspor, dan hasil analisis. Verifikator memeriksa koordinat. Viewer memiliki akses baca pada panel sesuai RLS dan menu yang tersedia.",
        295: "Kebutuhan negatif mencakup pengecualian status unknown dari K-Means, larangan mengubah lokasi exact melalui edit biasa, larangan impor menetapkan exact, batas K terhadap jumlah titik, pembatasan tipe dan ukuran foto, kerahasiaan isi pengajuan, serta review yang hanya dapat dilakukan satu kali melalui transaksi terkunci.",
        298: "Kebutuhan non-fungsional meliputi integritas, keamanan, privasi, keterlacakan, responsivitas, aksesibilitas interaksi, dan reproduksibilitas. RLS serta RPC melindungi data; bucket foto privat dan signed URL membatasi bukti; audit dan snapshot menjaga jejak; paginasi dan marker clustering mengendalikan ukuran tampilan.",
        299: "Tidak semua kebutuhan dapat dibuktikan dari kode. Responsivitas, keyboard, browser, sesi autentikasi, kebijakan Storage, dan transaksi review memerlukan uji runtime. Lint dan build menguji konsistensi statis serta kemampuan menghasilkan bundle, tetapi tidak menggantikan pengujian integrasi Supabase.",
        303: "Lapisan presentasi menggunakan React dan Leaflet. Service layer mengakses Supabase PostgreSQL, Auth, RPC, dan Storage. Utility layer menjalankan normalisasi lokasi dan K-Means. Data persisten meliputi UMKM, profil, run analisis, batch impor, audit, pengajuan, pemulihan kode, dan objek foto.",
        306: "Arsitektur memisahkan presentation, service, utility, dan data layer. RootApp memuat aplikasi publik atau AdminApp secara lazy. BusinessSubmissionForm mengelola pengajuan serta pelacakan; halaman admin mengelola data, antrean, verifikasi, analisis, impor, dan audit; umkmService membungkus kueri, upload, signed URL, dan RPC.",
        307: "Jalur publik membaca hanya UMKM aktif dan published, tetapi dapat melakukan insert pengajuan, upload foto ke folder terbatas, memanggil RPC status berdasarkan UUID, serta membuat permintaan pemulihan. Jalur admin membaca data sesuai role dan memakai RPC untuk operasi sensitif seperti verifikasi, impor, serta persetujuan atau penolakan pengajuan.",
        312: "Entitas operasional berpusat pada umkm. Tabel umkm_submissions menampung data kiriman, tracking_code, referensi foto, status review, dan approved_umkm_id. Tabel umkm_submission_recovery_requests menampung permintaan pemulihan kode. Foto disimpan pada bucket privat umkm-submission-photos, bukan di dalam baris PostgreSQL.",
        313: "admin_profiles menyimpan role; kmeans_runs menyimpan hasil analisis; import_batches merekam impor; dan audit_logs merekam perubahan UMKM. Relasi reviewed_by, handled_by, verified_by, dan created_by mengaitkan proses dengan akun, sementara approved_umkm_id menghubungkan pengajuan yang disetujui ke record UMKM hasil transaksi.",
        314: "Constraint menjaga pasangan koordinat, rentang nilai, metadata exact, panjang snapshot, state pending/approved/rejected, pasangan photo_path dan photo_kind, serta state pemulihan. Indeks mendukung antrean menurut status dan waktu. RLS dan fungsi security definer membatasi read/write sesuai aktor.",
        319: "Pengguna publik dapat membaca peta, membuat pengajuan, mengunggah satu foto, melacak status, dan meminta pemulihan kode tanpa akun, tetapi tidak dapat membaca tabel pengajuan. Superadmin dan admin dapat mengelola UMKM, meninjau pengajuan, melihat foto, memproses pemulihan, mengimpor data, dan menyimpan run. Verifikator dapat memperbarui lokasi melalui flow verifikasi, sedangkan viewer hanya membaca informasi admin yang diizinkan.",
        322: "Lokasi perkiraan dapat dianalisis secara eksploratif, sedangkan status belum terverifikasi dikecualikan. Pengajuan yang disetujui tanpa titik menjadi belum terverifikasi; jika titik dikirim, status awalnya tetap perkiraan_kecamatan. Status tepat hanya diberikan setelah verifikator menjalankan RPC verifikasi dan konfirmasi manual.",
        327: "Bukti verifikasi yang dianjurkan meliputi sumber rujukan, tanggal pemeriksaan, dan screenshot. Foto pengajuan dapat membantu peninjauan identitas usaha, tetapi tidak otomatis membuktikan ketepatan koordinat dan tidak menggantikan konfirmasi melalui alur Verifikasi Lokasi.",
        342: "Snapshot lokal tanggal 10 September 2026 memuat 1.802 record: 1.743 berstatus perkiraan_kecamatan, 59 belum_terverifikasi, dan belum ada lokasi berstatus tepat. Karena isMappableLocation mengecualikan status unknown, eksperimen lokal menggunakan 1.743 titik dan mencatat 59 record sebagai excluded.",
        343: "Analisis sensitivitas hanya dapat membandingkan seluruh titik mappable dengan lokasi tepat setelah proses verifikasi menghasilkan sampel exact yang memadai. Pada snapshot saat ini perbandingan tersebut belum sah karena jumlah exact masih nol; kondisi ini dilaporkan sebagai keterbatasan data.",
        359: "Pengujian black-box mencakup peta dan pencarian, koleksi, UMKM terdekat, rute, share, pengajuan dengan foto/titik, pelacakan dan pemulihan, login, dashboard, CRUD, review, verifikasi, impor, ekspor, audit, K-Means, dan kontrol role. Setiap test case mencatat prasyarat, input, hasil aktual, status, serta bukti.",
        360: "Pengujian negatif mencakup koordinat tidak lengkap atau di luar rentang, tipe/ukuran foto tidak sah, tracking code tidak valid, review ulang, akses foto tanpa hak, import exact, verifikasi tanpa konfirmasi, viewer melakukan mutasi, dan K melebihi jumlah titik. Database, service, atau UI harus menolak operasi pada batas yang sesuai.",
        365: "",
        369: "Repositori memisahkan aplikasi publik dan panel admin. Halaman publik menggunakan App, Map, BusinessList, Sidebar, MapInfoPanel, CommunityCollections, dan BusinessSubmissionForm. Panel admin menyediakan Dashboard, Businesses, Verification, KMeans, ImportExport, Submissions, dan Audit.",
        370: "RootApp memilih jalur publik atau /admin dan memuat AdminApp secara lazy. Publik dapat mengeksplorasi data tanpa akun serta mengirim pengajuan ketika Supabase aktif. Operasi administrasi memerlukan session dan profil aktif, sementara kebijakan database tetap menjadi batas keamanan utama.",
        371: "Struktur kode memisahkan komponen, utility, dan service. kmeans.js menjalankan analisis, location.js membedakan koordinat analisis dan tampilan, publicDataService memilih database atau JSON lokal, dan umkmService mengelola operasi Supabase termasuk upload foto, signed URL, status pelacakan, review, serta dashboard overview.",
        372: "Pada pemuatan publik, sistem menampilkan state loading atau error, membentuk kategori, menerapkan filter, menghitung statistik kualitas, menyaring titik mappable, dan menjalankan K-Means sesuai K aktif. useMemo membatasi komputasi ulang, sedangkan markercluster dimuat dinamis untuk mengurangi beban tampilan peta.",
        374: "4.2 Implementasi Layanan Publik dan Peta",
        375: "Peta Leaflet menampilkan marker UMKM, centroid, area zona, popup, dan marker cluster. Pencarian memeriksa nama, merek, pemilik, alamat, dan kategori. Filter dapat berasal dari satu kategori atau koleksi komunitas, sedangkan pemilihan UMKM disimpan pada parameter ?umkm= agar tautan dapat dibagikan.",
        376: "Fitur UMKM terdekat meminta izin geolokasi setelah pengguna menekan tombol, menghitung Haversine terhadap koordinat tampilan, mengurutkan jarak, dan menampilkan lima hasil. Rute Google Maps juga memakai titik tampilan atau alamat sebagai fallback. Pilihan ini menjaga jarak dan tujuan publik konsisten dengan marker yang terlihat, sedangkan analisis K-Means tetap memakai analysis_lat dan analysis_lng.",
        377: "Formulir publik menerima nama usaha, pemilik, WhatsApp, kategori, alamat, catatan, satu foto opsional, dan titik opsional. Titik dapat dipilih pada peta atau melalui pencarian Nominatim yang dibatasi ke Sulawesi Utara, diberi jeda permintaan, dan disimpan dalam cache sesi. Foto dibatasi pada JPEG, PNG, atau WebP dengan ukuran maksimum 5 MB.",
        378: "Setelah insert berhasil, aplikasi menampilkan UUID pelacakan. RPC get_umkm_submission_status hanya mengembalikan nama usaha, status, catatan review, dan waktu, sehingga pemohon tidak memperoleh akses baca umum ke tabel. Jika kode hilang, nama usaha, pemilik, dan WhatsApp dikirim sebagai permintaan pemulihan untuk diverifikasi admin di luar kanal publik.",
        379: "Admin melihat antrean pengajuan tertua lebih dulu, foto melalui signed URL satu jam, titik kiriman, dan catatan. review_umkm_submission mengunci baris agar hanya pending yang dapat diproses. Persetujuan membuat record UMKM dan memperbarui approved_umkm_id dalam satu transaksi; titik kiriman tetap berstatus perkiraan, sedangkan pengajuan tanpa titik menjadi belum terverifikasi.",
        383: "Migrasi Supabase membangun admin_profiles, umkm, kmeans_runs, import_batches, audit_logs, umkm_submissions, dan umkm_submission_recovery_requests. Migrasi juga membuat bucket foto privat, indeks antrean, fungsi status pelacakan, fungsi review atomik, serta trigger pemulihan.",
        384: "Role superadmin, admin, verifikator, dan viewer diperiksa di UI dan database. Publik hanya membaca UMKM aktif-published, memasukkan pengajuan atau permintaan pemulihan, mengunggah foto ke folder submissions, dan menjalankan RPC status. Admin berwenang membaca pengajuan serta foto dan melakukan review.",
        387: "RLS memakai primitive is_admin, can_manage_umkm, dan can_verify_umkm. Kebijakan tabel serta Storage membatasi operasi per aktor, sedangkan fungsi security definer memeriksa role kembali. Dengan demikian, menyembunyikan menu atau tombol bukan satu-satunya lapisan keamanan.",
        388: "AuditPage menampilkan 100 perubahan UMKM terbaru. Dashboard memuat ringkasan run K-Means terakhir, antrean lima pengajuan tertua, antrean lima pemulihan tertua, indikator masalah koordinat/kategori, dan peringatan jika snapshot analisis berbeda dari data aktif saat ini.",
        389: "Impor dijalankan atomik melalui RPC dan menolak status exact dari CSV. Review pengajuan juga atomik: baris dikunci, status pending diverifikasi, record UMKM dibuat bila disetujui, lalu relasi persetujuan diperbarui. Kedua alur mencegah state setengah selesai pada operasi multi-langkah.",
        392: "Preprocessing memetakan alamat lama ke pusat wilayah dan membuat display offset deterministik agar marker yang bertumpuk tetap dapat dipilih. Snapshot JSON memuat 1.802 record dengan 90 nilai area dan 14 kode produk. Koordinat dasar tersebut merupakan estimasi administratif, bukan GPS usaha.",
        395: "location.js memprioritaskan analysis_lat/lng untuk K-Means dan display_lat/lng untuk marker, rute, serta UMKM terdekat. Data lama tanpa location_accuracy diberi default perkiraan_kecamatan. Pemisahan ini mencegah offset visual memengaruhi hasil clustering.",
        396: "isMappableLocation mensyaratkan status bukan belum_terverifikasi dan pasangan koordinat analisis tersedia. Pada snapshot lokal, 1.743 record memenuhi syarat dan 59 record dikecualikan meskipun masih mempunyai koordinat fallback, karena statusnya menyatakan lokasi belum dapat dipercaya untuk analisis.",
        399: "Statistik snapshot menunjukkan 96,73% record berstatus perkiraan dan 3,27% belum terverifikasi, tanpa lokasi exact. Hasil K-Means pada tahap ini harus dibaca sebagai pola koordinat wilayah. Dashboard dan filter kualitas membantu mengarahkan pekerjaan verifikasi sebelum analisis final digunakan untuk keputusan yang lebih rinci.",
        418: "Traceability menghubungkan peta dan penemuan ke App/Map/BusinessList; pengajuan dan pelacakan ke BusinessSubmissionForm, umkmService, Storage, dan migrasi; dashboard ke DashboardPage; kualitas lokasi ke location.js, VerificationPage, RPC, dan constraint; K-Means ke kmeans.js serta KMeansPage; dan audit/impor ke tabel serta halaman terkait.",
        419: "Matriks membuktikan ketersediaan artefak, bukan keberhasilan runtime. Pemeriksaan lint, build, dan eksekusi algoritma dapat dinyatakan sebagai hasil karena benar-benar dijalankan, sedangkan autentikasi, RLS, Storage, transaksi review, responsivitas, dan alur pengguna tetap memerlukan bukti black-box pada environment final.",
        421: "Pemeriksaan teknis pada 10 September 2026 menjalankan ESLint dan build Vite terhadap working tree proyek; keduanya selesai dengan exit code 0. Build mentransformasi 1.845 modul dan menghasilkan bundle produksi. Vite memberi peringatan bahwa chunk utama 623,84 kB sebelum gzip melebihi ambang 500 kB serta dynamic import umkmService tidak memisahkan modul karena service yang sama juga diimpor statis oleh panel admin.",
        423: "Pengujian algoritma menjalankan K=2 sampai K=10 pada 1.743 titik mappable dari JSON lokal. Pengulangan untuk K=3, K=5, dan K=8 menghasilkan WCSS, centroid, dan komposisi anggota yang identik, sesuai tujuan seed deterministik. Hash snapshot adalah 596d3b7e2785d53805e15b48fe7545cdde8095e37c588c5b29d7869ba982e2f3.",
        424: "Pengujian black-box terintegrasi belum dinyatakan lulus karena memerlukan session Supabase, role uji, data uji terpisah, kebijakan Storage aktif, dan browser pada beberapa viewport. Status pada tabel membedakan hasil yang benar-benar dieksekusi dari fitur yang baru didukung artefak kode dan migrasi.",
        427: "Eksekusi snapshot lokal menghasilkan WCSS untuk K=2 sampai K=10 sebagaimana Tabel 4.7. Nilai turun tajam pada beberapa transisi, tetapi K=9 menghasilkan WCSS lebih tinggi daripada K=8. Secara teoritis optimum global tidak meningkat ketika K bertambah; kenaikan ini menunjukkan solusi lokal dari satu inisialisasi deterministik per K dan menegaskan bahwa calculateElbowData belum menjalankan multi-start.",
        428: "Tabel 4.7 Hasil Eksekusi K-Means pada Snapshot Data Lokal",
        429: "Dataset eksperimen berisi 1.743 titik, seluruhnya berstatus perkiraan_kecamatan. Nilai WCSS, iterasi, dan komposisi klaster dihitung langsung oleh performKMeans pada kode proyek. K=2 menghasilkan 3.678.351,83 km², sedangkan K=10 menghasilkan 114.319,76 km², tetapi urutan penurunannya tidak cukup stabil untuk menetapkan satu nilai K optimal.",
        430: "Penurunan terbesar terjadi pada K=3 terhadap K=2 sebesar 66,85%, diikuti K=5 terhadap K=4 sebesar 60,12%. Namun penurunan lain berfluktuasi dan K=9 naik 10,53% terhadap K=8. Karena itu elbow tunggal tidak dipilih dari tabel ini. Evaluasi lanjutan perlu menjalankan beberapa inisialisasi per K atau menyimpan solusi terbaik, lalu membandingkan stabilitas klaster.",
        431: "Halaman publik menghitung cluster berdasarkan filter interaktif, sedangkan run admin menggunakan data aktif, published, dan mappable. Tabel hasil menggunakan snapshot JSON lokal yang eksplisit agar dapat direproduksi dan tidak disamakan dengan state filter pengguna atau data produksi yang mungkin telah berubah.",
        432: "Dominasi koordinat perkiraan dapat mengecilkan jarak dalam satu kecamatan dan mengubah posisi centroid. Oleh karena itu hasil numerik dilaporkan bersama komposisi kualitas lokasi: 1.743 perkiraan, 59 unknown yang dikecualikan, dan nol exact. Nilai K final ditunda sampai data lebih presisi atau prosedur multi-start diterapkan.",
        434: "Kekuatan sistem terletak pada penggabungan pemetaan, partisipasi publik, dan tata kelola data. Pemohon dapat mengajukan usaha serta memantau status, tetapi publik tidak diberi akses membaca tabel pengajuan. Admin meninjau foto dan data melalui URL sementara, sedangkan persetujuan, verifikasi, serta audit dijaga oleh database.",
        448: "Dari sisi usability, publik memperoleh pencarian, koleksi, deep-link, share, rute, dan UMKM terdekat. Pengajuan dan pelacakan memakai dialog dengan focus restoration dan Escape handler. Admin memperoleh dashboard berbasis antrean serta filter kualitas. Keberadaan atribut aksesibilitas dan layout responsif tetap perlu divalidasi melalui keyboard, screen reader, dan perangkat nyata.",
        455: "Keterbatasan keempat adalah cakupan pengujian runtime. Lint, build, dan algoritma lokal telah dijalankan, tetapi autentikasi, RLS, Storage, review atomik, pemulihan kode, dan perilaku lintas perangkat belum didokumentasikan sebagai black-box. Build juga menunjukkan chunk utama yang besar sehingga optimasi code-splitting layak menjadi pekerjaan lanjutan.",
        462: "1. Sistem Web GIS telah mengintegrasikan peta, pencarian, filter, koleksi, UMKM terdekat, rute, share, pengajuan publik, pelacakan, serta panel admin berbasis React, Leaflet, dan Supabase.",
        463: "2. Skema membedakan lokasi tepat, perkiraan, dan belum terverifikasi serta melindungi pengajuan, foto, review, dan pemulihan kode melalui RLS, bucket privat, RPC, constraint, dan trigger.",
        464: "3. Snapshot lokal memuat 1.802 UMKM; 1.743 titik perkiraan masuk analisis dan 59 titik berstatus belum terverifikasi dikecualikan. Belum ada lokasi exact pada snapshot tersebut.",
        465: "4. K-Means++ deterministik dengan Haversine menghasilkan output yang sama pada pengulangan K=3, K=5, dan K=8, tetapi deret WCSS K=2 sampai K=10 tidak monoton pada K=9. Nilai K optimal belum dapat ditetapkan dari satu run per K.",
        466: "5. ESLint dan build produksi berhasil pada 10 September 2026. Pengujian integrasi Supabase dan black-box lintas perangkat masih harus dilaksanakan dan dilampirkan sebelum klaim kelulusan seluruh fungsi dibuat.",
        467: "Secara keseluruhan, proyek menghubungkan representasi spasial, partisipasi pemilik usaha, kualitas koordinat, dan analisis clustering. Setiap alur mempunyai batas data yang berbeda: marker dan rute memakai koordinat tampilan, K-Means memakai koordinat analisis, dan lokasi exact hanya lahir dari verifikasi.",
        468: "Snapshot, hash, parameter, WCSS, centroid, radius, dan riwayat run mendukung keterlacakan. Temuan WCSS K=9 memperlihatkan bahwa determinisme tidak sama dengan optimalitas; prosedur ilmiah perlu menambah multi-start atau membandingkan solusi sebelum memilih elbow.",
        469: "Dari sisi governance, publik dapat berpartisipasi tanpa memperoleh akses ke data pengajuan lain. Admin meninjau bukti melalui signed URL dan transaksi persetujuan menghubungkan pengajuan dengan record UMKM. Desain ini memperluas sistem dari peta statis menjadi workflow data yang dapat diaudit.",
        471: "1. Meningkatkan verifikasi koordinat lapangan agar proporsi lokasi exact cukup untuk analisis sensitivitas.",
        472: "2. Menjalankan K-Means dengan beberapa inisialisasi per K dan menyimpan solusi WCSS terbaik sebelum memilih elbow.",
        473: "3. Melaksanakan pengujian integrasi untuk Auth, RLS, Storage, review, pelacakan, dan pemulihan kode menggunakan akun serta data uji terpisah.",
        474: "4. Menguji aksesibilitas dan responsivitas pada keyboard, screen reader, browser, serta perangkat bergerak.",
        475: "5. Mengoptimalkan code-splitting agar chunk utama produksi tidak melampaui ambang peringatan build.",
        476: "6. Menambahkan ekspor peta dan laporan hasil klaster serta kebijakan retensi foto pengajuan yang ditolak atau kedaluwarsa.",
        477: "",
        439: "",
        498: "",
        503: "Repositori terdiri atas scripts untuk preprocessing dan web-app untuk aplikasi React. Bagian publik memuat App, Map, BusinessList, Sidebar, MapInfoPanel, CommunityCollections, dan BusinessSubmissionForm. Bagian admin memuat Dashboard, Businesses, Verification, KMeans, ImportExport, Submissions, dan Audit. Service mengakses Supabase, utility menangani lokasi/K-Means/CSV, dan empat migrasi membangun fondasi admin, pengajuan, foto-pelacakan, serta pemulihan kode.",
    }
    for index, text in replacements.items():
        document.paragraphs[index].text = text
    document.paragraphs[53].alignment = WD_ALIGN_PARAGRAPH.LEFT

    page_numbers = {
        58: 1, 59: 1, 60: 5, 61: 6, 62: 7, 63: 7, 64: 8,
        65: 10, 66: 10, 67: 11, 68: 12, 69: 13, 70: 15, 71: 16,
        72: 17, 73: 18, 74: 18, 75: 20, 76: 21, 77: 23,
        78: 24, 79: 24, 80: 24, 81: 26, 82: 27, 83: 27, 84: 28,
        85: 30, 86: 32, 87: 33, 88: 34, 89: 35, 90: 36, 91: 36,
        92: 36, 93: 37, 94: 38, 95: 38, 96: 40,
        97: 41, 98: 41, 99: 42, 100: 43, 101: 45, 102: 46,
        103: 48, 104: 49, 105: 50, 106: 52, 107: 53, 108: 57,
        109: 57, 110: 59, 111: 59, 112: 60, 113: 62,
        114: 64, 115: 64, 116: 64, 117: 64,
        120: 26, 121: 29, 122: 29, 123: 30, 124: 32, 125: 33,
        126: 34, 127: 39, 128: 43, 129: 45, 130: 48,
        132: 20, 133: 21, 134: 21, 135: 22, 136: 27, 137: 27,
        138: 30, 139: 31, 140: 32, 141: 37, 142: 38, 143: 38,
        144: 39, 145: 40, 146: 41, 147: 44, 148: 49, 149: 49,
        150: 50, 151: 51, 152: 52, 153: 56, 154: 64, 155: 64,
    }
    for index, page_number in page_numbers.items():
        title = document.paragraphs[index].text.rsplit("\t", 1)[0]
        document.paragraphs[index].text = f"{title}\t{page_number}"

    toc_start = document.paragraphs[58]
    front_matter = [
        ("HALAMAN JUDUL", "i"),
        ("HALAMAN PENGESAHAN", "iii"),
        ("PERNYATAAN KEASLIAN TULISAN", "iv"),
        ("ABSTRAK", "v"),
        ("ABSTRACT", "vi"),
        ("KATA PENGANTAR", "vii"),
        ("DAFTAR ISI", "ix"),
        ("DAFTAR GAMBAR", "xi"),
        ("DAFTAR TABEL", "xii"),
    ]
    for title, page_number in front_matter:
        toc_item = deepcopy(toc_start._p)
        item = Paragraph(toc_item, toc_start._parent)
        item.text = f"{title}\t{page_number}"
        toc_start._p.addprevious(toc_item)


def capture_table_format(table) -> list[list[tuple[object | None, object | None]]]:
    templates = []
    for row_index in range(min(2, len(table.rows))):
        row_templates = []
        for cell in table.rows[row_index].cells:
            paragraph = cell.paragraphs[0]
            run = paragraph.runs[0] if paragraph.runs else None
            row_templates.append((deepcopy(paragraph._p.pPr), deepcopy(run._r.rPr) if run is not None else None))
        templates.append(row_templates)
    return templates


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def set_table(table, rows: list[list[str]]) -> None:
    formats = capture_table_format(table)
    while len(table.rows) < len(rows):
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.rows[row_index].cells[column_index]
            cell.text = value
            paragraph = cell.paragraphs[0]
            template_row = 0 if row_index == 0 else min(1, len(formats) - 1)
            template_column = min(column_index, len(formats[template_row]) - 1)
            p_pr, r_pr = formats[template_row][template_column]
            if p_pr is not None:
                current = paragraph._p.pPr
                if current is not None:
                    paragraph._p.remove(current)
                paragraph._p.insert(0, deepcopy(p_pr))
            run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
            if r_pr is not None:
                current = run._r.rPr
                if current is not None:
                    run._r.remove(current)
                run._r.insert(0, deepcopy(r_pr))
    set_repeat_header(table.rows[0])


def update_tables(document: Document) -> None:
    tables = document.tables
    set_table(tables[2], [
        ["Teknologi", "Lapisan", "Peran"],
        ["React 19", "Presentation", "Komponen, state, dialog, dan interaksi UI"],
        ["Vite 8", "Build tooling", "Development server, lint workflow, dan build produksi"],
        ["Leaflet 1.9 / react-leaflet 5", "Web GIS", "Peta, marker, popup, zona, dan pemilih lokasi"],
        ["leaflet.markercluster 1.5", "Visualisasi", "Pengelompokan marker sesuai zoom"],
        ["Supabase JS 2.112.3", "Service client", "Auth, query, RPC, Storage, dan signed URL"],
        ["PostgreSQL / Supabase", "Data dan keamanan", "Tabel, constraint, trigger, transaksi, dan RLS"],
        ["Supabase Storage", "Bukti pengajuan", "Penyimpanan foto privat maksimal 5 MB"],
        ["Nominatim OpenStreetMap", "Pencarian lokasi", "Mencari jalan atau gedung dalam viewbox Sulawesi Utara"],
        ["Web Crypto SHA-256", "Reproduksibilitas", "Hash snapshot input K-Means"],
    ])
    set_table(tables[4], [
        ["Konsep", "Prinsip", "Implementasi pada Sistem"],
        ["SIG Web", "Atribut terhubung dengan posisi", "Leaflet, daftar, popup, pencarian, filter, rute, dan UMKM terdekat"],
        ["Partisipasi publik", "Data baru masuk melalui tinjauan", "Pengajuan, foto, tracking_code, review, dan pemulihan"],
        ["Kualitas lokasi", "Ketelitian dinyatakan eksplisit", "tepat, perkiraan_kecamatan, belum_terverifikasi"],
        ["K-Means++", "Seeding dan assignment centroid", "Seed deterministik dengan jarak Haversine"],
        ["Reproduksibilitas", "Input dan parameter terlacak", "input_snapshot, dataset_hash, parameters, kmeans_runs"],
        ["Data governance", "Mutasi dan bukti dibatasi", "RBAC, RPC, RLS, bucket privat, constraint, audit"],
    ])
    set_table(tables[6], [
        ["Kode", "Kebutuhan Fungsional"],
        ["F-01", "Menampilkan UMKM pada peta, popup, zona, dan daftar interaktif."],
        ["F-02", "Mencari, memfilter, memilih koleksi, membuka rute, membagikan tautan, dan menemukan lima UMKM terdekat."],
        ["F-03", "Menerima pengajuan UMKM dengan titik dan satu foto opsional."],
        ["F-04", "Memberikan kode pelacakan, status review, dan kanal pemulihan kode."],
        ["F-05", "Mengautentikasi pengguna admin dan membatasi menu serta operasi sesuai role."],
        ["F-06", "Menambah, mengubah, memublikasikan, mengaktifkan, atau menonaktifkan UMKM."],
        ["F-07", "Meninjau pengajuan; persetujuan membuat UMKM secara atomik, penolakan menyimpan alasan."],
        ["F-08", "Memverifikasi koordinat tepat melalui alur yang mencatat verifier dan waktu."],
        ["F-09", "Membedakan lokasi tepat, perkiraan, dan belum terverifikasi."],
        ["F-10", "Menjalankan K-Means untuk nilai K yang dipilih."],
        ["F-11", "Menampilkan centroid, anggota, WCSS, iterasi, radius, dan ringkasan zona."],
        ["F-12", "Menyimpan run K-Means beserta parameter, hash, dan snapshot."],
        ["F-13", "Mendukung dashboard pekerjaan, filter kualitas, impor/ekspor, dan audit."],
    ])
    set_table(tables[7], [
        ["Kode", "Kebutuhan Non-Fungsional"],
        ["NF-01", "Aplikasi dapat digunakan pada browser modern, desktop, dan perangkat bergerak."],
        ["NF-02", "Data administratif dilindungi autentikasi, role, RPC, dan Row Level Security."],
        ["NF-03", "Foto pengajuan disimpan pada bucket privat dan dibaca admin melalui signed URL sementara."],
        ["NF-04", "Constraint menjaga pasangan koordinat, rentang, state review, dan metadata verifikasi."],
        ["NF-05", "Hasil analisis dapat dilacak melalui parameter, SHA-256, dan snapshot input."],
        ["NF-06", "Marker clustering, memoization, paginasi, dan lazy loading menjaga beban antarmuka."],
        ["NF-07", "Impor dan review penting bersifat atomik; perubahan UMKM dicatat pada audit log."],
        ["NF-08", "RPC pelacakan hanya mengembalikan data status minimum tanpa membuka tabel pengajuan."],
    ])
    set_table(tables[8], [
        ["Entitas", "Fungsi Utama"],
        ["umkm", "Atribut usaha, koordinat analisis/tampilan, kualitas lokasi, status aktif/published, dan metadata verifikasi."],
        ["admin_profiles", "Profil pengguna dan role superadmin, admin, verifikator, atau viewer."],
        ["kmeans_runs", "Parameter, hasil, statistik, hash, dan snapshot setiap eksekusi clustering."],
        ["import_batches", "Metadata proses impor atomik dan jumlah baris."],
        ["audit_logs", "Perubahan INSERT, UPDATE, atau DELETE pada UMKM."],
        ["umkm_submissions", "Data pengajuan, koordinat opsional, foto, kode pelacakan, review, dan relasi UMKM hasil persetujuan."],
        ["umkm_submission_recovery_requests", "Permintaan pemulihan kode serta metadata penyelesaian."],
        ["storage.objects", "Objek foto pada bucket privat umkm-submission-photos."],
    ])
    set_table(tables[10], [
        ["Fungsi", "Superadmin", "Admin", "Verifikator", "Viewer"],
        ["Melihat dashboard dan data", "Ya", "Ya", "Ya", "Ya"],
        ["Tambah/edit/nonaktif UMKM", "Ya", "Ya", "Tidak", "Tidak"],
        ["Verifikasi lokasi", "Ya", "Ya", "Ya", "Tidak"],
        ["Jalankan K-Means", "Ya", "Ya", "Ya", "Ya"],
        ["Simpan hasil K-Means", "Ya", "Ya", "Tidak", "Tidak"],
        ["Impor/ekspor administratif", "Ya", "Ya", "Tidak", "Tidak"],
        ["Tinjau pengajuan, foto, dan pemulihan", "Ya", "Ya", "Tidak", "Tidak"],
        ["Melihat riwayat audit", "Ya", "Ya", "Ya", "Ya"],
    ])
    set_table(tables[13], [
        ["ID", "Skenario", "Hasil yang Diharapkan"],
        ["T-01", "Membuka halaman publik", "Peta dan daftar tampil tanpa akun admin."],
        ["T-02", "Pencarian, filter, koleksi, dan deep-link", "Hasil dan pilihan UMKM konsisten dengan input/URL."],
        ["T-03", "UMKM terdekat, rute, dan share", "Lima hasil memakai titik tampilan; tautan dapat dibuka/dibagikan."],
        ["T-04", "Mengirim pengajuan dengan/ tanpa titik dan foto", "Validasi berjalan dan UUID pelacakan diterima."],
        ["T-05", "Melacak status dan meminta pemulihan", "Status minimum tampil; permintaan pemulihan tercatat."],
        ["T-06", "Login admin valid/tidak valid", "Akses hanya diberikan pada session dan profil aktif."],
        ["T-07", "Dashboard dan filter kualitas", "Antrean, masalah data, dan ringkasan run dapat dibuka."],
        ["T-08", "CRUD dan status aktif/published", "Perubahan tersimpan sesuai hak akses."],
        ["T-09", "Setujui atau tolak pengajuan", "Keputusan satu kali; persetujuan membuat UMKM secara atomik."],
        ["T-10", "Verifikasi lokasi tepat", "Koordinat dan metadata verifier tersimpan konsisten."],
        ["T-11", "Mengubah exact di luar flow verifikasi", "Database menolak perubahan sensitif."],
        ["T-12", "K-Means K=2 sampai K=10", "WCSS, iterasi, centroid, radius, dan anggota dihasilkan."],
        ["T-13", "Mengulangi K dengan snapshot sama", "WCSS, centroid, dan komposisi identik."],
        ["T-14", "Menyimpan run K-Means", "Hash, snapshot, parameter, dan hasil tersimpan."],
        ["T-15", "Impor CSV dan ekspor dataset", "Validasi, transaksi, dan file keluaran berjalan sesuai aturan."],
        ["T-16", "Menjalankan ESLint", "Pemeriksaan selesai tanpa error."],
        ["T-17", "Menjalankan build produksi", "Bundle produksi terbentuk dan peringatan dicatat."],
    ])
    set_table(tables[14], [
        ["ID", "Skenario", "Bukti"],
        ["T-18", "Filter koordinat/kategori/lokasi perlu diperiksa", "Screenshot tabel dan jumlah hasil"],
        ["T-19", "Viewer mencoba operasi mutasi", "Pesan UI dan penolakan RLS/RPC"],
        ["T-20", "Verifikator menyimpan exact tanpa konfirmasi", "Pesan penolakan"],
        ["T-21", "Koordinat tidak lengkap atau di luar rentang", "Pesan validasi/constraint"],
        ["T-22", "Impor mencoba status exact", "Pesan penolakan service/RPC"],
        ["T-23", "Impor valid dijalankan atomik", "Batch sukses dan jumlah record"],
        ["T-24", "Foto salah tipe atau lebih dari 5 MB", "Pesan validasi dan tidak ada objek tersimpan"],
        ["T-25", "Akses foto tanpa role manage", "Penolakan kebijakan Storage"],
        ["T-26", "Kode pelacakan salah atau tidak ditemukan", "Pesan aman tanpa membuka data lain"],
        ["T-27", "Pemulihan ditandai selesai", "handled_at dan handled_by terisi"],
        ["T-28", "Record unknown masuk kandidat K-Means", "excluded_records bertambah"],
        ["T-29", "Perubahan UMKM", "Entri audit_logs"],
        ["T-30", "Desktop/mobile dan navigasi keyboard", "Screenshot, rekaman langkah, dan audit aksesibilitas"],
    ])
    set_table(tables[16], [
        ["Modul", "Tanggung Jawab"],
        ["App.jsx", "State publik, filter, zonasi, deep-link, share, dan pembukaan dialog pengajuan/pelacakan."],
        ["Map.jsx", "Marker, centroid/area, UMKM terdekat, rute, dan koordinat tampilan."],
        ["BusinessList / CommunityCollections", "Daftar, koleksi kategori, pencarian, pemilihan, dan detail UMKM."],
        ["BusinessSubmissionForm.jsx", "Pengajuan, foto, pencarian/pilih lokasi, tracking, dan pemulihan kode."],
        ["AdminApp.jsx", "Session, profile role, routing, lazy load, dan proteksi menu."],
        ["DashboardPage.jsx", "Antrean kerja, kualitas data, ringkasan zona, dan deteksi snapshot berubah."],
        ["BusinessesPage.jsx", "CRUD, arsip, paginasi, dan filter kualitas/status."],
        ["VerificationPage.jsx", "Pemeriksaan koordinat dan status akurasi."],
        ["SubmissionsPage.jsx", "Review, signed URL foto, kode pelacakan, dan pemulihan."],
        ["KMeansPage.jsx", "Run, hash snapshot, statistik, infografis, dan riwayat."],
        ["ImportExportPage.jsx", "Validasi CSV, impor atomik, dan ekspor aman spreadsheet."],
        ["AuditPage.jsx", "Penyajian 100 perubahan UMKM terbaru."],
        ["services / utils", "Akses Supabase, normalisasi lokasi, K-Means, dan CSV."],
    ])
    set_table(tables[17], [
        ["Kontrol", "Lokasi Implementasi", "Fungsi"],
        ["Rentang dan pasangan koordinat", "Service + CHECK", "Mencegah koordinat tidak sah/setengah lengkap"],
        ["Exact metadata invariant", "CHECK + trigger", "Mewajibkan koordinat dan verifier"],
        ["RBAC", "admin_profiles + fungsi akses", "Membedakan manage, verify, dan read"],
        ["Verifikasi lokasi", "RPC verify_umkm_location", "Membatasi perubahan lokasi sensitif"],
        ["Pengajuan publik", "RLS insert + constraint", "Hanya menerima state pending yang valid"],
        ["Review pengajuan", "RPC review_umkm_submission", "Mengunci baris dan membuat UMKM secara atomik"],
        ["Privasi foto", "Bucket privat + policy + signed URL", "Membatasi upload dan pembacaan bukti"],
        ["Pelacakan", "RPC get_umkm_submission_status", "Mengembalikan data status minimum berdasar UUID"],
        ["Pemulihan kode", "RLS + trigger handler", "Mencatat penyelesaian dan admin pemroses"],
        ["Impor atomik", "RPC import_umkm_batch", "Menghindari partial import"],
        ["Audit dan reproduksibilitas", "audit_logs + kmeans_runs", "Mencatat mutasi serta input analisis"],
    ])
    set_table(tables[19], [
        ["Kebutuhan", "Status berdasarkan Artefak", "Catatan"],
        ["Peta dan daftar publik", "Tersedia", "App, Map, BusinessList, Sidebar."],
        ["Pencarian, koleksi, nearby, rute, share", "Tersedia", "App, Map, BusinessList, CommunityCollections."],
        ["Pengajuan, foto, pelacakan, pemulihan", "Tersedia", "BusinessSubmissionForm, service, Storage, tiga migrasi."],
        ["Dashboard dan filter kualitas", "Tersedia", "DashboardPage dan BusinessesPage."],
        ["Panel admin dan RBAC", "Tersedia", "AdminApp, admin_profiles, RLS/RPC."],
        ["Review pengajuan atomik", "Tersedia", "SubmissionsPage dan review_umkm_submission."],
        ["Verifikasi koordinat", "Tersedia", "VerificationPage, service, RPC, constraint."],
        ["K-Means, WCSS, radius", "Tersedia", "kmeans.js dan KMeansPage."],
        ["Hash, snapshot, riwayat", "Tersedia", "SHA-256 dan kmeans_runs."],
        ["Impor/ekspor dan audit", "Tersedia", "CSV utility, RPC import, audit_logs."],
        ["Pengujian runtime final", "Belum lengkap", "Lint/build/algoritma telah dijalankan; integrasi dan black-box tertunda."],
    ])
    set_table(tables[20], [
        ["ID", "Fungsi", "Status pada 10 September 2026"],
        ["T-01", "Halaman publik", "Artefak dan build tersedia; black-box final tertunda"],
        ["T-02", "Pencarian/filter/koleksi", "Didukung kode; black-box final tertunda"],
        ["T-03", "Nearby/rute/share", "Didukung kode; black-box final tertunda"],
        ["T-04", "Pengajuan dan foto", "Didukung kode/migrasi; integrasi Storage tertunda"],
        ["T-05", "Pelacakan/pemulihan", "Didukung RPC/migrasi; integrasi tertunda"],
        ["T-06", "Login dan role", "Didukung Auth/RLS; pengujian akun tertunda"],
        ["T-07", "Dashboard/filter kualitas", "Didukung kode; black-box final tertunda"],
        ["T-08", "CRUD UMKM", "Didukung service/RLS; integrasi tertunda"],
        ["T-09", "Review pengajuan", "Didukung RPC atomik; integrasi tertunda"],
        ["T-10", "Verifikasi koordinat", "Didukung RPC/constraint; integrasi tertunda"],
        ["T-11", "Proteksi lokasi exact", "Didukung trigger/constraint; uji negatif tertunda"],
        ["T-12", "K-Means K=2..10", "Lulus pada snapshot lokal 1.743 titik"],
        ["T-13", "Determinisme K=3, 5, 8", "Lulus; hasil pengulangan identik"],
        ["T-14", "Penyimpanan run", "Didukung skema/service; integrasi tertunda"],
        ["T-15", "Impor/ekspor", "Didukung kode/RPC; uji file tertunda"],
        ["T-16", "ESLint", "Lulus, exit code 0"],
        ["T-17", "Build produksi", "Lulus, exit code 0; ada peringatan ukuran chunk"],
    ])
    set_table(tables[21], [
        ["Kelompok", "Jumlah Skenario", "Status Saat Ini", "Bukti"],
        ["Peta publik dan penemuan", "3", "Perlu black-box", "Artefak + build"],
        ["Pengajuan, foto, pelacakan", "3", "Perlu integrasi", "Kode + SQL"],
        ["Autentikasi dan role", "3", "Perlu akun uji", "RLS/RPC"],
        ["CRUD, review, verifikasi", "5", "Perlu integrasi", "Service + SQL"],
        ["Impor, ekspor, dan audit", "4", "Perlu pengujian file/DB", "Kode + RPC"],
        ["K-Means dan reproduksibilitas", "4", "Sebagian lulus", "Output + hash"],
        ["Lint dan build", "2", "Lulus", "Exit code 0"],
        ["Responsif dan aksesibilitas", "1", "Perlu perangkat", "Screenshot/audit"],
    ])
    set_table(tables[22], [
        ["K", "WCSS km²", "Perubahan terhadap K sebelumnya", "Catatan"],
        ["2", "3.678.351,83", "-", "Baseline; 623 dan 1.120 anggota"],
        ["3", "1.219.410,63", "Turun 66,85%", "Penurunan besar"],
        ["4", "1.044.680,59", "Turun 14,33%", "Perbaikan relatif kecil"],
        ["5", "416.619,60", "Turun 60,12%", "Penurunan besar kedua"],
        ["6", "253.603,88", "Turun 39,13%", "Perbaikan masih nyata"],
        ["7", "226.310,51", "Turun 10,76%", "Perbaikan kecil"],
        ["8", "142.468,24", "Turun 37,05%", "Penurunan kembali besar"],
        ["9", "157.472,06", "Naik 10,53%", "Solusi lokal; kurva tidak monoton"],
        ["10", "114.319,76", "Turun 27,40%", "Belum menetapkan K optimal"],
    ])


def replace_diagrams(document: Document, files: dict[str, Path]) -> None:
    replaced = set()
    for relationship in document.part.rels.values():
        if not relationship.reltype.endswith("/image"):
            continue
        name = Path(relationship.target_ref).name
        if name in files:
            relationship.target_part._blob = files[name].read_bytes()
            replaced.add(name)
    missing = set(files) - replaced
    if missing:
        raise RuntimeError(f"Image parts not found: {sorted(missing)}")


def resize_diagrams(document: Document) -> None:
    sizes = {
        "image7.png": (6.1, 2.58),
        "image9.png": (6.0, 1.2),
    }
    for shape in document.inline_shapes:
        blips = shape._inline.xpath(".//a:blip")
        if not blips:
            continue
        relationship_id = blips[0].get(qn("r:embed"))
        image_name = Path(document.part.rels[relationship_id].target_ref).name
        if image_name not in sizes:
            continue
        width, height = sizes[image_name]
        shape.width = Inches(width)
        shape.height = Inches(height)


def main() -> None:
    document = Document(SOURCE)
    replace_paragraphs(document)
    update_tables(document)
    replace_diagrams(document, diagram_files())
    resize_diagrams(document)
    document.core_properties.title = "Implementasi Algoritma K-Means Clustering pada Web GIS untuk Pemetaan dan Analisis Zonasi Persebaran UMKM di Sulawesi Utara"
    document.core_properties.subject = "Revisi isi skripsi sesuai implementasi proyek gis-umkm-mapping"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
