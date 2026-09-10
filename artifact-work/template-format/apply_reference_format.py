from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path(r"D:\projek UMKM\gis-umkm-mapping-main\Skripsi_GIS_diperbarui_sesuai_proyek_pembaruan.docx")
OUTPUT = Path(r"D:\projek UMKM\gis-umkm-mapping-main\Skripsi_GIS_format_mengikuti_BISINDO.docx")


def set_font(paragraph: object, bold: bool | None = None) -> None:
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        run.font.size = Pt(12)
        if bold is not None:
            run.bold = bold


def set_body_format(paragraph: object, first_line: bool = True) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(6)
    fmt.first_line_indent = Inches(0.5) if first_line else None
    set_font(paragraph)


def set_front_heading(paragraph: object) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = Inches(0.25)
    fmt.first_line_indent = None
    set_font(paragraph, bold=True)


def set_toc_entry(paragraph: object, level: int) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.left_indent = Inches({1: 0, 2: 0.18, 3: 0.36}[level])
    fmt.first_line_indent = None
    fmt.tab_stops.clear_all()
    fmt.tab_stops.add_tab_stop(Inches(5.51), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    set_font(paragraph)


def toc_level(text: str) -> int:
    label = text.split("\t", 1)[0]
    if re.match(r"^\d+\.\d+\.\d+\s", label):
        return 3
    if re.match(r"^\d+\.\d+\s", label) or label.startswith("Lampiran "):
        return 2
    return 1


def replace_paragraph_text(paragraph: object, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def configure_kata_pengantar(doc: Document) -> None:
    title = doc.paragraphs[51]
    opening = doc.paragraphs[52]
    lead_in = doc.paragraphs[53]
    closing = doc.paragraphs[54]
    signature = doc.paragraphs[55]

    replace_paragraph_text(title, "KATA PENGANTAR")
    set_front_heading(title)

    replace_paragraph_text(
        opening,
        "Puji syukur dipanjatkan ke hadirat Tuhan Yang Maha Esa atas rahmat dan karunia-Nya sehingga skripsi ini dapat diselesaikan. Skripsi ini disusun untuk melengkapi salah satu syarat kelulusan Program Sarjana Terapan pada Program Studi D-IV Teknik Informatika, Jurusan Teknik Elektro, Politeknik Negeri Manado.",
    )
    set_body_format(opening)

    replace_paragraph_text(
        lead_in,
        "Penulis menyadari bahwa tanpa bantuan dan bimbingan dari berbagai pihak, sangatlah sulit bagi penulis untuk menyelesaikan skripsi ini. Oleh karena itu, penulis menyampaikan ucapan terima kasih kepada:",
    )
    set_body_format(lead_in)

    acknowledgements = [
        "Direktur Politeknik Negeri Manado;",
        "Ketua Jurusan Teknik Elektro;",
        "Koordinator Program Studi D-IV Teknik Informatika;",
        "Ketua Pelaksana Ujian Skripsi beserta seluruh panitia;",
        "Dosen Pembimbing I dan Dosen Pembimbing II atas arahan, bimbingan, dan masukan selama penyusunan skripsi ini;",
        "Seluruh dosen dan staf pengajar Program Studi D-IV Teknik Informatika atas ilmu, arahan, dan dukungan akademik;",
        "YBLI Yayasan Bina Lentera Insan selaku mitra penelitian yang telah mendukung pelaksanaan penelitian ini;",
        "Orang tua, keluarga, sahabat, dan rekan-rekan yang telah membantu penulis menyelesaikan skripsi ini.",
    ]
    for number, text in enumerate(acknowledgements, start=1):
        paragraph = closing.insert_paragraph_before(f"{number}. {text}")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fmt = paragraph.paragraph_format
        fmt.left_indent = Inches(0.31)
        fmt.first_line_indent = Inches(-0.31)
        fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(1)
        set_font(paragraph)

    replace_paragraph_text(
        closing,
        "Akhir kata, penulis menyadari skripsi ini masih jauh dari sempurna. Semoga skripsi ini bermanfaat bagi pembaca, khususnya di bidang Teknik Informatika dan Sistem Informasi Geografis.",
    )
    set_body_format(closing)

    replace_paragraph_text(signature, "Manado, Agustus 2026")
    signature.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = signature.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = None
    set_font(signature)
    signature_index = next(index for index, paragraph in enumerate(doc.paragraphs) if paragraph._p is signature._p)
    page_break = doc.paragraphs[signature_index + 1]
    penulis = page_break.insert_paragraph_before("Penulis,")
    nama = page_break.insert_paragraph_before("\nPUTRA FEBRIAN")
    for paragraph in (penulis, nama):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.first_line_indent = None
        set_font(paragraph)
    if nama.runs:
        nama.runs[-1].bold = True


def configure_toc(doc: Document) -> None:
    toc_index = next(index for index, paragraph in enumerate(doc.paragraphs) if paragraph.text == "DAFTAR ISI")
    list_end = next(
        index
        for index, paragraph in enumerate(doc.paragraphs[toc_index + 1 :], start=toc_index + 1)
        if paragraph.text == "DAFTAR GAMBAR"
    )
    set_front_heading(doc.paragraphs[toc_index])
    for paragraph in doc.paragraphs[toc_index + 1 : list_end]:
        if paragraph.text.strip():
            set_toc_entry(paragraph, toc_level(paragraph.text))

    front_pages = {
        "DAFTAR ISI\tix": "DAFTAR ISI\tviii",
        "DAFTAR GAMBAR\txi": "DAFTAR GAMBAR\txi",
        "DAFTAR TABEL\txii": "DAFTAR TABEL\txii",
    }
    for paragraph in doc.paragraphs[toc_index + 1 : list_end]:
        if paragraph.text in front_pages:
            replace_paragraph_text(paragraph, front_pages[paragraph.text])
            set_toc_entry(paragraph, 1)


def normalize_body(doc: Document) -> None:
    list_marker = re.compile(r"^(\d+[.)]|[a-z][.)])\s")
    for index, paragraph in enumerate(doc.paragraphs[166:], start=166):
        text = paragraph.text.strip()
        if not text or paragraph.style.name != "Normal" or len(text) < 50:
            continue
        if list_marker.match(text) or text.startswith(("Gambar ", "Tabel ", "Sumber:")):
            continue
        set_body_format(paragraph)


def main() -> None:
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    configure_kata_pengantar(doc)
    configure_toc(doc)
    normalize_body(doc)
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
