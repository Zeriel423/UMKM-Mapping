from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn


def length(value: object) -> float | None:
    return round(value.inches, 3) if value is not None else None


def paragraph_data(paragraph: object, index: int) -> dict[str, object]:
    fmt = paragraph.paragraph_format
    return {
        "index": index,
        "text": paragraph.text,
        "style": paragraph.style.name,
        "alignment": str(paragraph.alignment),
        "left_indent": length(fmt.left_indent),
        "right_indent": length(fmt.right_indent),
        "first_line_indent": length(fmt.first_line_indent),
        "space_before": length(fmt.space_before),
        "space_after": length(fmt.space_after),
        "line_spacing": str(fmt.line_spacing),
        "line_spacing_rule": str(fmt.line_spacing_rule),
        "keep_with_next": fmt.keep_with_next,
        "keep_together": fmt.keep_together,
        "runs": [
            {
                "text": run.text,
                "font": run.font.name,
                "size": run.font.size.pt if run.font.size else None,
                "bold": run.bold,
                "italic": run.italic,
            }
            for run in paragraph.runs
        ],
    }


def table_data(table: object, index: int) -> dict[str, object]:
    grid = table._tbl.tblGrid.gridCol_lst
    return {
        "index": index,
        "style": table.style.name if table.style else None,
        "rows": len(table.rows),
        "cols": len(table.columns),
        "grid_inches": [round(int(col.get(qn("w:w"))) / 1440, 3) for col in grid],
        "sample": [
            [cell.text.replace("\n", " | ") for cell in row.cells]
            for row in table.rows[:5]
        ],
    }


def inspect(path: Path) -> dict[str, object]:
    doc = Document(path)
    paragraphs = [paragraph_data(p, i) for i, p in enumerate(doc.paragraphs)]
    front_matter = [p for p in paragraphs if p["index"] < 130]
    return {
        "path": str(path),
        "sections": [
            {
                "index": i + 1,
                "width": length(section.page_width),
                "height": length(section.page_height),
                "top_margin": length(section.top_margin),
                "bottom_margin": length(section.bottom_margin),
                "left_margin": length(section.left_margin),
                "right_margin": length(section.right_margin),
                "header_distance": length(section.header_distance),
                "footer_distance": length(section.footer_distance),
                "different_first_page": section.different_first_page_header_footer,
            }
            for i, section in enumerate(doc.sections)
        ],
        "front_matter": front_matter,
        "tables": [table_data(table, i) for i, table in enumerate(doc.tables)],
    }


if __name__ == "__main__":
    output = inspect(Path(sys.argv[1]))
    Path(sys.argv[2]).write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
